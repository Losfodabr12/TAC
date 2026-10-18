from __future__ import annotations

import hashlib
import re
import unicodedata
import zipfile
from pathlib import Path

from html.parser import HTMLParser

SUPPORTED = {".txt", ".md", ".markdown", ".pdf", ".docx", ".html", ".htm"}


class TextHTML(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts, self.links, self.title_parts = [], [], []
        self.reference_links = []
        self.stack = []
        self.regions = {"content": [], "article": [], "main": []}

    def append(self, text):
        if self.stack and self.stack[-1][1]:
            return
        self.parts.append(text)
        active = set().union(*(item[2] for item in self.stack)) if self.stack else set()
        for region in active:
            self.regions[region].append(text)

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        classes = set(attrs.get("class", "").split())
        hidden = ((self.stack[-1][1] if self.stack else False)
                  or tag in ("script", "style", "noscript", "svg", "template", "nav", "footer", "aside")
                  or attrs.get("aria-hidden") == "true" or "hidden" in attrs
                  or bool(classes & {"mw-editsection", "mw-jump-link", "mw-indicators", "toc", "noprint", "vector-header-container", "vector-page-toolbar"}))
        region = set()
        if attrs.get("id") in ("mw-content-text", "article-body") or "mw-parser-output" in classes:
            region.add("content")
        if tag == "article":
            region.add("article")
        if tag == "main" or attrs.get("role") == "main":
            region.add("main")
        void = tag in ("br", "hr", "img", "meta", "link", "input", "source", "wbr", "area", "base", "col", "embed", "param", "track")
        if not void:
            reference = ((self.stack[-1][3] if self.stack else False)
                         or attrs.get('id', '').startswith('cite_note')
                         or 'reference-text' in classes or tag == 'cite')
            self.stack.append((tag, hidden, region, reference))
        if tag == "a" and not hidden:
            href = attrs.get("href")
            if href:
                self.links.append(href)
                if self.stack and self.stack[-1][3]:
                    self.reference_links.append(href)
        if not hidden and tag in ("p", "div", "br", "li", "section", "article", "tr", "h1", "h2", "h3", "h4"):
            self.append("\n")
        if not hidden and tag in ("h1", "h2", "h3", "h4"):
            self.append("#" * int(tag[-1]) + " ")
        if not hidden and tag in ("td", "th"):
            self.append(" ")

    def handle_endtag(self, tag):
        if tag in ("p", "div", "li", "h1", "h2", "h3", "h4", "tr"):
            self.append("\n")
        for index in range(len(self.stack) - 1, -1, -1):
            if self.stack[index][0] == tag:
                del self.stack[index:]
                break

    def handle_data(self, data):
        if any(item[0] == "title" for item in self.stack):
            self.title_parts.append(data)
        else:
            self.append(data)

    def text(self):
        for name in ("content", "article", "main"):
            text = normalize("".join(self.regions[name]))
            if len(text) >= 40:
                return text
        return normalize("".join(self.parts))


def normalize(text):
    text = unicodedata.normalize("NFC", text).replace("\x00", "")
    return "\n".join(re.sub(r"[\t\r ]+", " ", line).strip()
                     for line in text.splitlines() if line.strip()).strip()


def decode_text(raw):
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")


def html_text(raw):
    parser = TextHTML()
    parser.feed(decode_text(raw) if isinstance(raw, bytes) else raw)
    return normalize(" ".join(parser.title_parts)), parser.text(), parser.links


def html_references(raw):
    parser = TextHTML()
    parser.feed(decode_text(raw) if isinstance(raw, bytes) else raw)
    return list(dict.fromkeys(parser.reference_links))[:300]


def extract(path: Path, config, check=lambda: None):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"Formato não suportado: {suffix}")
    if path.stat().st_size > config.get("max_file_mb") * 2**20:
        raise ValueError("Arquivo maior que o limite configurado.")
    limit = config.get("max_text_chars")
    check()
    title, meta = path.stem, {"filename": path.name, "format": suffix}
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("PDF protegido. Importe uma cópia textual autorizada.")
        pieces, size = [], 0
        for page_num, page in enumerate(reader.pages):
            check()
            # Streams PDF descompactados também têm limite; OCR não é executado.
            if page.get_contents() is not None and len(page.get_contents().get_data()) > 16 * 2**20:
                raise ValueError("Página PDF excede o limite de descompactação.")
            text = page.extract_text() or ""
            size += len(text)
            if size > limit:
                raise ValueError("Texto excede o limite. Divida o documento.")
            pieces.append(f"\nPágina {page_num + 1}\n{text}")
        text = "\n".join(pieces)
        meta["pages"] = len(reader.pages)
        if not any(len(p.strip()) > 30 for p in pieces):
            raise ValueError("PDF sem texto extraível; OCR não está disponível nesta versão.")
    elif suffix == ".docx":
        with zipfile.ZipFile(path) as z:
            if sum(i.file_size for i in z.infolist()) > 40 * 2**20:
                raise ValueError("DOCX excede o limite de descompactação.")
        from docx import Document
        doc = Document(path)
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            check()
            lines.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        text = "\n".join(lines)
    elif suffix in (".html", ".htm"):
        parsed_title, text, _ = html_text(path.read_bytes())
        title = parsed_title or title
    else:
        text = decode_text(path.read_bytes())
    text = normalize(text)
    if not text:
        raise ValueError("Documento vazio ou sem texto extraível.")
    if len(text) > limit:
        raise ValueError("Texto excede o limite configurado. Divida o documento.")
    check()
    return title, text, meta


CHUNKING_VERSION = 2


def heading_level(line):
    """Falas, nomes seguidos de dois pontos e linhas curtas não são títulos."""
    if re.match(r"^#{1,6}\s+\S", line) or re.match(r"^Página \d+\b", line):
        return 2
    plain = "".join(c for c in unicodedata.normalize("NFKD", line) if not unicodedata.combining(c))
    plain = re.sub(r"^[^\w]+", "", plain).strip()
    if re.match(r"^(?:EPISODIO|CAPITULO|PARTE|LIVRO)\s+(?:\d+|[IVXLCDM]+)\b", plain, re.I):
        return 2
    if line.startswith(("—", "–", "- ", '"', "“")) or line.endswith(":"):
        return 0
    letters = [c for c in line if c.isalpha()]
    if len(line) <= 100 and len(letters) >= 4 and all(c.isupper() for c in letters):
        return 1
    return 0


def split_sections(text, size=1200):
    """Trechos contíguos, sem sobreposição, preservando contexto de diálogos."""
    section, buffer, length = "Início", [], 0
    for line in text.splitlines():
        heading = heading_level(line)
        boundary = heading == 2 or (heading == 1 and length >= size // 3)
        if (boundary or length + len(line) > size) and buffer:
            yield section, "\n".join(buffer)
            buffer, length = [], 0
        if heading and not buffer:
            section = line.lstrip("# ")[:160]
        for start in range(0, len(line), size):
            part = line[start:start + size]
            if length + len(part) > size and buffer:
                yield section, "\n".join(buffer)
                buffer, length = [], 0
            buffer.append(part)
            length += len(part) + 1
    if buffer:
        yield section, "\n".join(buffer)


def text_hash(text):
    return hashlib.sha256(normalize(text).encode("utf-8")).hexdigest()
