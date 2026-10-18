import zipfile

import pytest

from tac.documents import extract, html_text


def test_html_no_script_execution(core, tmp_path):
    path = tmp_path / "a.html"
    path.write_text("<html><title>Fonte local</title><script>delete_everything()</script><h1>Aprendizado</h1><p>Texto &amp; referência.</p></html>", encoding="utf-8")
    title, text, meta = extract(path, core.config)
    assert title == "Fonte local"
    assert "Texto & referência." in text
    assert "delete_everything" not in text


def test_html_prioritizes_article_and_ignores_navigation():
    paragraph = "Brasil é um país da América do Sul. Este texto pertence ao artigo e contém informação documental verificável na própria fonte."
    title, text, links = html_text(f"<title>Brasil</title><nav>Menu principal</nav><main><div>Barra de ferramentas</div><div id='mw-content-text'><h2>Introdução</h2><p>{paragraph}</p></div></main><footer>Outros sites</footer>")
    assert title == "Brasil"
    assert paragraph in text
    assert "Menu principal" not in text
    assert "Barra de ferramentas" not in text
    assert "Outros sites" not in text
    assert "## Introdução" in text


def test_docx_extract(core, tmp_path):
    from docx import Document
    path = tmp_path / "a.docx"
    doc = Document()
    doc.add_heading("Título de teste", 1)
    doc.add_paragraph("Informação verificável em DOCX.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Tabela textual"
    doc.save(path)
    _, text, _ = extract(path, core.config)
    assert "Informação verificável" in text
    assert "Tabela textual" in text


def test_pdf_text(core, tmp_path):
    # PDF mínimo textual produzido no próprio teste, sem depender de modelos/fontes externos.
    path = tmp_path / "a.pdf"
    stream = b"BT /F1 12 Tf 50 700 Td (Conteudo PDF verificavel para a TAC.) Tj ET"
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>", b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>", b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"]
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(str(i).encode() + b" 0 obj\n" + obj + b"\nendobj\n")
    xref = len(data)
    data.extend(b"xref\n0 6\n0000000000 65535 f \n")
    for value in offsets[1:]:
        data.extend(f"{value:010d} 00000 n \n".encode())
    data.extend(f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    path.write_bytes(data)
    _, text, meta = extract(path, core.config)
    assert "Conteudo PDF verificavel" in text
    assert meta["pages"] == 1


def test_bad_document_does_not_stop_queue(core, tmp_path):
    bad = tmp_path / "invalido.pdf"
    bad.write_bytes(b"not a pdf")
    good = tmp_path / "valido.md"
    good.write_text("# Documento válido\nA aplicação continua processando depois de uma falha isolada.", encoding="utf-8")
    core.enqueue_file(bad)
    core.enqueue_file(good)
    core.gate.set_enabled(True)
    result = core.perform("Estudo", core.run_queue, learning=True)
    assert result["processed"] == 1
    assert [r["status"] for r in core.db.query("SELECT * FROM queue ORDER BY id")] == ["erro", "concluido"]
    assert len(core.memory.sources()) == 1


def test_size_limits(core, tmp_path):
    path = tmp_path / "large.txt"
    core.config.update(max_file_mb=1)
    path.write_bytes(b"a" * (2**20 + 1))
    with pytest.raises(ValueError, match="limite"):
        core.enqueue_file(path)
